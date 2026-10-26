# -*- coding: utf-8 -*-
"""
md_to_docx_manuscript.py — manuscript_en.md → manuscript_en.docx（BMC/SNAPP 上传版）

转换范围：标题层级、粗体/斜体、段落、Unicode 上下标（原字符保留）、
编号声明列表（按字面编号，不启用 Word 自动编号，避免编号漂移）。
不转换表格（正文无 markdown 表格；附表以 CSV 作为 additional files 另传）。

内置验证（审计纪律：数字不得在转换中损坏）：
1) 纯文本词数：docx 提取文本 vs md 去语法文本，偏差 < 1%；
2) 关键 token 抽查：审计敏感数字与结论短语逐一在 docx 中存在。
"""
import os
import re

from docx import Document
from docx.shared import Pt

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE, "05_稿件", "manuscript_en.md")
OUT = os.path.join(BASE, "05_稿件", "manuscript_en.docx")

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*)")


def add_runs(par, text):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            par.add_run(tok[1:-1]).italic = True
        else:
            par.add_run(tok)


def strip_md(text):
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*([^*\n]+?)\*", r"\1", text)
    text = re.sub(r"^#{1,6}\s+", "", text)
    text = re.sub(r"^>\s?", "", text, flags=re.M)
    return text


def main():
    with open(SRC, encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for ln in lines:
        s = ln.rstrip()
        if not s.strip():
            continue
        if s == "---":
            continue
        if s.startswith("# "):
            doc.add_heading(s[2:].strip(), level=0)
        elif s.startswith("## "):
            doc.add_heading(s[3:].strip(), level=1)
        elif s.startswith("### "):
            doc.add_heading(s[4:].strip(), level=2)
        else:
            p = doc.add_paragraph()
            add_runs(p, s.strip())

    doc.save(OUT)

    # ---------------- 验证 ----------------
    d = Document(OUT)
    docx_text = "\n".join(p.text for p in d.paragraphs)
    md_text = strip_md("\n".join(l for l in lines
                                 if l.strip() and l.strip() != "---"))
    wc_docx = len(docx_text.split())
    wc_md = len(md_text.split())
    dev = abs(wc_docx - wc_md) / max(wc_md, 1)
    assert dev < 0.01, f"词数偏差 {dev:.2%} 超阈（docx {wc_docx} vs md {wc_md}）"

    probes = [
        "2.8 × 10⁻⁶", "0.71–1.77", "26/26", "MPC2 −0.45", "PGAM1 +1.25",
        "1/C(21,10)", "REML", "AutoDock Vina 1.2.7", "metafor 5.2.1",
        "not executed", "did not measure Kpy itself",
        "AI-assisted technology disclosure", "[6,36]", "GSE245703",
        "The Second Affiliated Hospital of Fujian University of Traditional Chinese Medicine",
        "0009-0004-9698-6552",
    ]
    missing = [t for t in probes if t not in docx_text]
    assert not missing, f"关键 token 在 docx 中缺失: {missing}"

    n_para = len([p for p in d.paragraphs if p.text.strip()])
    print(f"OK  -> {OUT}")
    print(f"    段落 {n_para} | 词数 docx {wc_docx} vs md {wc_md}（偏差 {dev:.2%}）")
    print(f"    关键 token 抽查 {len(probes)}/{len(probes)} 通过（含审计敏感数字）")


if __name__ == "__main__":
    main()
